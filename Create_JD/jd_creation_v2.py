from langgraph.graph import START,StateGraph,END
from typing import Dict, Any, List
import os
import tempfile
from pathlib import Path
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.runnables import RunnableParallel,RunnableLambda
from langchain_text_splitters import RecursiveCharacterTextSplitter
from qdrant_client import QdrantClient
from langchain_qdrant import QdrantVectorStore
from operator import itemgetter
from Shared.prompts import jd_creation_prompt,jd_judge_scores_prompt,jd_outline_prompt
from Shared.state import JDCreationState
from Shared.config import jd_creation_emb_llm,jd_creation_judge_llm,jd_creation_llm_1,jd_creation_llm_2,jd_creation_llm_3,jd_creation_outline_llm

DEFAULT_POLICY_PATH = str(Path(__file__).parent / "HR_Hiring_Policy.pdf")


def _load_policy_docs() -> List[Document]:
    """Return LangChain Documents from the org policy or the default PDF."""
    import logging
    logger = logging.getLogger(__name__)

    try:
        from database.session import get_session_factory, session_scope
        from database.repo import get_org_settings
        from services import s3_service

        with session_scope(get_session_factory()) as session:
            settings = get_org_settings(session)
            s3_key = settings.policy_s3_key
            filename = settings.policy_filename or ""
    except Exception as e:
        logger.warning("Could not read org policy settings: %s", e)
        s3_key = None
        filename = ""

    if s3_key:
        suffix = Path(filename).suffix.lower()
        try:
            local_path = s3_service.LOCAL_STORE / s3_key if s3_service._LOCAL_FALLBACK else None
            if local_path and local_path.exists():
                raw_path = str(local_path)
            else:
                import boto3, io
                client = boto3.client("s3", region_name=s3_service.REGION)
                buf = io.BytesIO()
                client.download_fileobj(s3_service.BUCKET, s3_key, buf)
                buf.seek(0)
                tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
                tmp.write(buf.read())
                tmp.flush()
                raw_path = tmp.name

            if suffix == ".pdf":
                return list(PyPDFLoader(raw_path).lazy_load())
            elif suffix == ".docx":
                from docx import Document as DocxDocument
                docx = DocxDocument(raw_path)
                text = "\n".join(p.text for p in docx.paragraphs if p.text.strip())
                return [Document(page_content=text, metadata={"source": filename})]
            else:
                logger.warning("Unsupported policy file type '%s'. Please upload a PDF or DOCX.", suffix)
        except Exception as e:
            logger.error("Failed to load custom policy '%s': %s", filename, e)

    return list(PyPDFLoader(DEFAULT_POLICY_PATH).lazy_load())

def build_query(x: Dict[str, Any]) -> str:
    """Construct a compact retrieval query from intake fields."""
    return (
        f"JD policy guidance for '{x.get('title','')}' in '{x.get('department','')}', "
        f"seniority '{x.get('seniority','')}', location '{x.get('location','')}', "
        f"employment '{x.get('employment_type','')}'. "
        f"Focus: responsibilities; must-have vs nice-to-have; skills/tools; "
        f"hiring stages & approvals; compensation/benefits phrasing; legal notes. "
        f"HR notes: {x.get('hr_notes','')}"
    )

class JD_Creation:
    def __init__(self,checkpointer):
        # Initialize all instance variables
        self.QDRANT_URL = os.environ["QDRANT_URL"]
        self.QDRANT_API_KEY = os.environ["QDRANT_API_KEY"]
        self.COLLECTION = os.getenv("QDRANT_COLLECTION", "test")
        self.client = QdrantClient(url=self.QDRANT_URL, api_key=self.QDRANT_API_KEY)
        
        self.emb_llm =jd_creation_emb_llm
        self.outline_llm =jd_creation_outline_llm
        self.jd_llm_1 = jd_creation_llm_1
        self.jd_llm_2 =jd_creation_llm_2
        self.jd_llm_3 =jd_creation_llm_3
        self.jd_judge_llm=jd_creation_judge_llm
        self.parser = StrOutputParser()

        graph = StateGraph(JDCreationState)
        graph.add_node('outline_creation', self.outline_creation)
        graph.add_node('jd_creation_1', self.jd_creation_1)
        graph.add_node('jd_creation_2', self.jd_creation_2)
        graph.add_node('jd_creation_3', self.jd_creation_3)
        graph.add_node('jd_judge', self.jd_judge)
        graph.add_node('jd_score', self.jd_score)
        graph.add_node('hr_approve', self.hr_approve)
        graph.add_node('await_hr', self.await_hr)

        graph.add_conditional_edges(
            START,
            lambda state: "skip_to_final" if state.get("hr_approved") else "outline_creation",
            {
                "skip_to_final": "hr_approve",
                "outline_creation": "outline_creation"
            }
        )
        graph.add_edge('outline_creation', 'jd_creation_1')
        graph.add_edge('outline_creation', 'jd_creation_2')
        graph.add_edge('outline_creation', 'jd_creation_3')
        graph.add_edge('jd_creation_1', 'jd_judge')
        graph.add_edge('jd_creation_2', 'jd_judge')
        graph.add_edge('jd_creation_3', 'jd_judge')
        graph.add_edge('jd_judge', 'jd_score')

        graph.add_conditional_edges(
            "jd_score",
            self.decision_router,
            {
                "refine": "outline_creation",
                "await": "await_hr",
                "approve": "hr_approve",
                "revise": "outline_creation",
            },
        )
        graph.add_edge('hr_approve', END)

        # Don't use nested checkpointer - let parent orchestrator handle it
        self.jd_workflow = graph.compile()

    def run(self, initial_state: dict) -> dict:
        """Run JD creation workflow and return final state dict."""
        
        # No need for thread_id config since we're not using nested checkpointing
        result = self.jd_workflow.invoke(initial_state)
        return result   
    
    def get_vectorstore(self, chunk):
        try:
            self.client.get_collection(self.COLLECTION)
            return QdrantVectorStore(
                client=self.client,
                collection_name=self.COLLECTION,
                embedding=self.emb_llm
            )
        except Exception:
            return QdrantVectorStore.from_documents(
                documents=chunk,
                embedding=self.emb_llm,
                url=self.QDRANT_URL,
                api_key=self.QDRANT_API_KEY,
                collection_name=self.COLLECTION
            )
        
    def formatted_docs(self, retrieved_doc):
        context = "\n\n".join(d.page_content for d in retrieved_doc)
        return context

    def outline_creation(self, state: JDCreationState):
        title = state['title']
        department = state["department"]
        seniority = state['seniority']
        location = state['location']
        employment_type = state['employment_type']
        hr_notes = state['hr_notes']
        feedback=state.get("review_feedback","")

        doc = _load_policy_docs()
        split = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=150)
        chunk = split.split_documents(doc)
        vectorstore = self.get_vectorstore(chunk)
        retriever = vectorstore.as_retriever(search_kwargs={'k': 6}, search_type='mmr')
        parallel_inputs = RunnableParallel({
                "policy_passages": RunnableLambda(build_query) | retriever | RunnableLambda(self.formatted_docs),
                "title": itemgetter("title"),
                "department": itemgetter("department"),
                "seniority": itemgetter("seniority"),
                "location": itemgetter("location"),
                "employment_type": itemgetter("employment_type"),
                "hr_notes": itemgetter("hr_notes"),
                "review_feedback": itemgetter("review_feedback"),   # <-- add this

            })

        outline_chain = parallel_inputs | jd_outline_prompt | self.outline_llm | self.parser

        result = outline_chain.invoke({
                "title": title,
                "department": department,
                "seniority": seniority,
                "location": location,
                "employment_type": employment_type,
                "hr_notes": hr_notes,
                'review_feedback':feedback,

            })
        return {
            'outline': result,
            "candidates": [],    # reset to avoid duplicates on revise cycle
            "judgements": {},
                }

    def jd_creation_1(self, state: JDCreationState):
        outline = state['outline']
        feedback=state.get("review_feedback","")
        
        jd_chain = jd_creation_prompt | self.jd_llm_1 | self.parser
        jd_text = jd_chain.invoke({'outline': outline,'review_feedback':feedback})
        return {"candidates": [{"id": "JD_A", "model_id": "jd_llm_1", "jd_text": jd_text}]}

    def jd_creation_2(self, state: JDCreationState):
        outline = state['outline']
        feedback=state.get("review_feedback","")

        jd_chain = jd_creation_prompt | self.jd_llm_2 | self.parser
        jd_text = jd_chain.invoke({'outline': outline,'review_feedback':feedback})
        return {"candidates": [{"id": "JD_B", "model_id": "jd_llm_2", "jd_text": jd_text}]}

    def jd_creation_3(self, state: JDCreationState):
        outline = state['outline']
        feedback=state.get("review_feedback","")
        
        jd_chain = jd_creation_prompt | self.jd_llm_3 | self.parser
        jd_text = jd_chain.invoke({'outline': outline,'review_feedback':feedback})
        return {"candidates": [{"id": "JD_C", "model_id": "jd_llm_3", "jd_text": jd_text}]}

    def jd_judge(self, state: JDCreationState):
        candidates = state["candidates"]
        review_feedback=state.get("review_feedback","")

        if len(candidates) < 3:
            return {
                "judgements": {},
                "status": "running"
            }
        A, B, C = candidates[0], candidates[1], candidates[2]

        outline = state['outline']

        msgs = jd_judge_scores_prompt.format_messages(
            title=state["title"],
            department=state["department"],
            seniority=state["seniority"],
            location=state["location"],
            employment_type=state["employment_type"],
            hr_notes=state["hr_notes"],
            outline=outline,
            review_feedback=review_feedback,
            model_id_A=A["model_id"], JD_A=A["jd_text"],
            model_id_B=B["model_id"], JD_B=B["jd_text"],
            model_id_C=C["model_id"], JD_C=C["jd_text"],
        )
        resp = self.jd_judge_llm.invoke(msgs)
        output_text = (resp.content or "").strip()
        
        judgements: Dict[str, Dict[str, Any]] = {}
        for raw_line in output_text.splitlines():
            line = raw_line.strip()
            if not line.startswith("JD_"):
                continue
            try:
                jd_id, rest = line.split(":", 1)
            # support en dash or hyphen, split once
                if "–" in rest:
                    score_str, feedback = rest.split("–", 1)
                elif "-" in rest:
                    score_str, feedback = rest.split("-", 1)
                else:
                    continue
                judgements[jd_id.strip()] = {
                    "score": float(score_str.strip().replace("/10","")),
                    "feedback": feedback.strip(),
                }
            except Exception:
                continue
        return {
                "judgements": judgements,   # <-- now a dict
            }

    def jd_score(self, state: JDCreationState):
        judgements = state.get("judgements") or {}
        if not judgements:
            return {}

        threshold = state["pass_threshold"]

        # pick highest scorer
        winner_id, winner_data = max(
            judgements.items(),
            key=lambda item: item[1].get("score", 0)
        )

        winner = {
            "id": winner_id,
            "score": winner_data.get("score", 0.0),
            "feedback": winner_data.get("feedback", ""),
            "jd_text": None,
            "model_id": None,
        }

        # attach JD text + model_id for the winner
        for cand in state.get("candidates", []):
            if cand.get("id") == winner_id:
                winner["jd_text"] = cand.get("jd_text")
                winner["model_id"] = cand.get("model_id")
                break

        # build feedback summary (used if we need to refine)
        outline_feedback = "\n".join(
            f"- {jid}: {data.get('score', 0)}/10 — {data.get('feedback','').strip()}"
            for jid, data in judgements.items()
        ).strip()

        passes_threshold = bool(winner["score"] >= threshold)

        updates = {
            "winner": winner,
            "passes_threshold": passes_threshold,
            "status": "await_hr" if passes_threshold else "running",
        }

        if not passes_threshold:
            updates.update({
                "outline_feedback": outline_feedback,
                "candidates": [],
                "judgements": {},
            })

        return updates

    def decision_router(self, state: JDCreationState) -> str:
        pt = state.get("passes_threshold")
        if pt is False:
            return "refine"     # failed score → regenerate via outline
        if pt is True:
            if state.get("skip_hr_wait"):
                return "approve"
            approved = state.get("hr_approved")
            if approved is True:
                return "approve"
            if approved is False:
                return "revise"  # HR rejected → back to outline (feedback already in state from API)
            return "await"      # show JD to HR and wait for input
        return "await"

    def hr_approve(self, state: JDCreationState):
        jd = state["winner"]
        iter = state["iteration"]
        approved = bool(state.get("hr_approved")) or bool(state.get("skip_hr_wait"))

        if approved:
            return {
                "final_jd": jd["jd_text"],
                "status": "done",
                "orchestration_status": "JD_Gen_Agent",
                "iteration": iter + 1,
            }
        return {
            "final_jd": None,
            "status": "needs_review",
        }

    def await_hr(self, state: JDCreationState):
        return {
            "status": "await_hr"
        }
